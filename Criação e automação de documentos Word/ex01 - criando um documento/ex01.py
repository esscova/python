"""
Criando um documento Word

Objetivo:
    Criar um documento word com a biblioteca python-docx

Autor:
    Wellington M Santos
"""

# BIBLIOTECAS
from docx import Document
from loguru import logger

# MAIN
def main():
    logger.info('Configurando a criação do documento word.')
    documento = Document()

    logger.info('Inserindo textos')
    documento.add_heading(text='hello world!', level=1)
    documento.add_paragraph(text='from python')

    logger.info('Salvando documento')
    documento.save('ex01.docx')
    logger.info('Documento criado com sucesso.')

if __name__ == '__main__':
    main()