"""
Formatação de textos e criação de listas

Objetivo:
    Aprendar a formatar palavras específicas (negrito/italico) e criar listas.
    

Autor:
    Wellington M Santos
"""
# BIBLIOTECAS
from docx import Document
from loguru import logger

# CONFIGURAÇÕES
BULLET  = 'List Bullet'
NUMBER = 'List Number'

# MAIN
def main():
    logger.info('Configurando a criação documento word')
    doc = Document()

    # titulo
    doc.add_heading('Ata da Reunião - Projeto Automação', 1)

    # paragrafo simples
    doc.add_paragraph('Reunião realizada em 15/08/2026 para alinhamento do status do projeto de automação de documentos.')

    # paragrafo com rusn
    paragrafo = doc.add_paragraph()
    paragrafo.add_run('Participantes: ').bold=True
    paragrafo.add_run('Wellington M. Santos, Equipe de Desenvolvimento e Stakeholders.\n')
    paragrafo.add_run('Objetivo principal: ').bold=True
    paragrafo.add_run('definir próximos passos e responsabilidades.').italic=True

    # subtitulo
    doc.add_heading('Pauta e Decisões', 2)

    # lista com marcadores
    doc.add_heading('ites discutidos: ', 3)
    doc.add_paragraph('Status atual da biblioteca python-docx nos scripts', style=BULLET)
    doc.add_paragraph('Padronização de templates de documentos', style=BULLET)
    doc.add_paragraph('Integração futura com geração de relatórios a partir de dados', style=BULLET)

    # lista numerada
    doc.add_heading('Ações definidas: ', 3)
    doc.add_paragraph('Criar uma série de exemplos progressivos (ex01 a ex0N)', style=NUMBER)
    doc.add_paragraph('Documentar padrões de logging e estrutura de arquivos', style=NUMBER)
    doc.add_paragraph('Preparar caso real de relatório automatizado', style=NUMBER)

    # paragrafo de citacao
    doc.add_paragraph('A automação de documentos reduz erros manuais e acelera entregas.', style='Intense Quote')

    # paragrafo final
    p_final = doc.add_paragraph()
    p_final.add_run('Próxima reunião: ').bold=True
    p_final.add_run('16/08/2026.\n')
    p_final.add_run('Responsável pelo registro: ').bold=True
    p_final.add_run('Wellington M. Santos')

    doc.save('ex02.docx')

if __name__ == '__main__':
    main()